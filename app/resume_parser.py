"""Resume document parsing utilities for HireMind AI.

Step 1: extract text from PDF and DOCX resumes and normalize it.
"""

from pathlib import Path
import re


SUPPORTED_EXTENSIONS = {".pdf", ".docx"}


def clean_text(text: str) -> str:
    """Normalize extracted resume text while preserving readable content."""
    text = text.replace("\x00", " ")
    text = re.sub(r"[ \t]+", " ", text)
    text = re.sub(r"\n{3,}", "\n\n", text)
    return text.strip()


def extract_from_pdf(file_path: str | Path) -> str:
    """Extract text from a PDF resume using PyMuPDF."""
    import fitz

    path = Path(file_path)
    if not path.is_file():
        raise FileNotFoundError(f"Resume not found: {path}")

    document = fitz.open(path)
    try:
        text = "\n".join(page.get_text("text") for page in document)
    finally:
        document.close()

    return clean_text(text)


def extract_from_docx(file_path: str | Path) -> str:
    """Extract paragraphs and table text from a DOCX resume."""
    from docx import Document

    path = Path(file_path)
    if not path.is_file():
        raise FileNotFoundError(f"Resume not found: {path}")

    document = Document(path)
    parts = [paragraph.text for paragraph in document.paragraphs if paragraph.text.strip()]

    for table in document.tables:
        for row in table.rows:
            cells = [cell.text.strip() for cell in row.cells if cell.text.strip()]
            if cells:
                parts.append(" | ".join(cells))

    return clean_text("\n".join(parts))


def extract_text(file_path: str | Path) -> str:
    """Extract and clean text from a supported resume file."""
    path = Path(file_path)
    extension = path.suffix.lower()

    if extension not in SUPPORTED_EXTENSIONS:
        supported = ", ".join(sorted(SUPPORTED_EXTENSIONS))
        raise ValueError(f"Unsupported file type '{extension}'. Supported types: {supported}")

    if extension == ".pdf":
        return extract_from_pdf(path)

    return extract_from_docx(path)


if __name__ == "__main__":
    import argparse

    parser = argparse.ArgumentParser(description="Extract text from a PDF or DOCX resume.")
    parser.add_argument("file", help="Path to the resume")
    args = parser.parse_args()

    print(extract_text(args.file))
